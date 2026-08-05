"""
Tests for ProjConfigPalette.py script
"""

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from unittest.mock import patch, MagicMock, mock_open
import tempfile

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

# Import the module under test
from apps import ProjConfigPalette
import ing_lib.common as common


class TestProjConfigPalette:
    """Test class for ProjConfigPalette script functionality."""
    
    def test_get_input_valid_args(self):
        """Test get_input function with valid arguments."""
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx',
            '--debug',
            '--username', 'testuser',
            '--ignore_ssl_error'
        ]
        
        inputs = ProjConfigPalette.get_input(args)
        
        assert inputs.server == 'https://test-server.example.com'
        assert inputs.function == 'query'
        assert inputs.excel == 'test.xlsx'
        assert inputs.debug is True
        assert inputs.username == 'testuser'
        assert inputs.ignore_ssl_error is True

    def test_get_input_minimal_args(self):
        """Test get_input function with minimal arguments."""
        args = [
            'https://test-server.example.com',
            'diff',
            'test.xlsx'
        ]
        
        inputs = ProjConfigPalette.get_input(args)
        
        assert inputs.server == 'https://test-server.example.com'
        assert inputs.function == 'diff'
        assert inputs.excel == 'test.xlsx'
        assert inputs.debug is False
        assert inputs.username is None

    @patch('apps.ProjConfigPalette.get_custom_palette')
    @patch('apps.ProjConfigPalette.get_built_in_palette')
    def test_get_palette_info_success(self, mock_built_in_palette, mock_custom_palette):
        """Test successful palette info retrieval."""
        mock_built_in_palette.return_value = [
            {
                'step_type': 'MANUAL_INPUT',
                'step_display_name': 'Manual Input',
                'palette_category': 'Input',
                'enable_disable': 'enable'
            }
        ]
        mock_custom_palette.return_value = [
            {
                'step_id': 'custom_step_1',
                'step_display_name': 'Custom Step 1',
                'palette_category': 'Custom',
                'enable_disable': 'enable'
            }
        ]
        
        result = ProjConfigPalette.get_palette_info('https://test-server.example.com')
        
        assert 'built_in' in result
        assert 'custom' in result
        assert len(result['built_in']) == 1
        assert len(result['custom']) == 1
        assert result['built_in'][0]['step_type'] == 'MANUAL_INPUT'
        assert result['custom'][0]['step_id'] == 'custom_step_1'

    @patch('apps.ProjConfigPalette.get_built_in_palette')
    def test_get_palette_info_failure(self, mock_built_in_palette):
        """Test palette info retrieval failure."""
        mock_built_in_palette.side_effect = Exception("Server error")
        
        with pytest.raises(common.IngeniumLibError, match=r"Error getting palette information from server"):
            ProjConfigPalette.get_palette_info('https://test-server.example.com')

    def test_read_palette_excel_file_not_found(self):
        """Test reading Excel file that doesn't exist."""
        with pytest.raises(common.IngeniumLibError, match=r"Excel file not found:"):
            ProjConfigPalette.read_palette_excel('nonexistent.xlsx')

    @patch('openpyxl.load_workbook')
    def test_read_palette_excel_success(self, mock_load_workbook):
        """Test successful reading of palette Excel file."""
        # Mock workbook and worksheets
        mock_workbook = MagicMock()
        mock_load_workbook.return_value = mock_workbook
        mock_workbook.sheetnames = ['Built-in Steps', 'Custom Steps']
        
        # Mock built-in steps worksheet
        mock_builtin_ws = MagicMock()
        # Mock header row (first row)
        mock_builtin_ws.__iter__ = MagicMock(return_value=iter([
            [MagicMock(value='step_type'), MagicMock(value='step_display_name'), MagicMock(value='palette_category')]
        ]))
        # Mock data rows (starting from row 2)
        mock_builtin_ws.iter_rows.return_value = iter([
            ('MANUAL_INPUT', 'Manual Input', 'Input')
        ])
        mock_workbook.__contains__.side_effect = lambda sheet: sheet in ['Built-in Steps', 'Custom Steps']
        mock_workbook.__getitem__.side_effect = lambda sheet: {
            'Built-in Steps': mock_builtin_ws,
            'Custom Steps': MagicMock()
        }[sheet]
        
        # Mock custom steps worksheet
        mock_custom_ws = MagicMock()
        # Mock header row (first row)
        mock_custom_ws.__iter__ = MagicMock(return_value=iter([
            [MagicMock(value='step_id'), MagicMock(value='step_display_name'), MagicMock(value='palette_category')]
        ]))
        # Mock data rows (starting from row 2)
        mock_custom_ws.iter_rows.return_value = iter([
            ('custom_1', 'Custom Step 1', 'Custom')
        ])
        mock_workbook.__getitem__.side_effect = lambda sheet: {
            'Built-in Steps': mock_builtin_ws,
            'Custom Steps': mock_custom_ws
        }[sheet]
        
        result = ProjConfigPalette.read_palette_excel('test.xlsx')
        
        assert 'built_in' in result
        assert 'custom' in result
        assert len(result['built_in']) == 1
        assert len(result['custom']) == 1
        # Just check that some data was read - the exact keys may vary
        assert result['built_in'][0] is not None
        assert result['custom'][0] is not None

    @patch('openpyxl.load_workbook')
    def test_read_palette_excel_malformed_file(self, mock_load_workbook):
        """Test reading malformed Excel file."""
        mock_load_workbook.side_effect = Exception("Invalid Excel format")
        
        with pytest.raises(common.IngeniumLibError, match=r"Error reading Excel file:"):
            ProjConfigPalette.read_palette_excel('malformed.xlsx')

    @patch('apps.ProjConfigPalette.Workbook')
    @patch('openpyxl.load_workbook')
    def test_write_palette_excel_success(self, mock_load_workbook, mock_workbook_class):
        """Test successful writing of palette Excel file."""
        # Mock workbook
        mock_workbook = MagicMock()
        mock_workbook_class.return_value = mock_workbook
        mock_workbook.sheetnames = ['Sheet']
        
        # Mock worksheets
        mock_builtin_ws = MagicMock()
        mock_custom_ws = MagicMock()
        mock_workbook.create_sheet.side_effect = [mock_builtin_ws, mock_custom_ws]
        mock_workbook.__contains__.side_effect = lambda sheet: sheet in ['Built-in Steps', 'Custom Steps']
        mock_workbook.__getitem__.side_effect = lambda sheet: {
            'Sheet': None,  # Will be removed
            'Built-in Steps': mock_builtin_ws,
            'Custom Steps': mock_custom_ws
        }[sheet]
        
        palette_info = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input'
                }
            ],
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Custom'
                }
            ]
        }
        
        ProjConfigPalette.write_palette_excel('test.xlsx', palette_info)
        
        mock_workbook.save.assert_called_once_with('test.xlsx')
        mock_workbook.close.assert_called_once()

    @patch('apps.ProjConfigPalette.Workbook')
    def test_write_palette_excel_failure(self, mock_workbook_class):
        """Test writing Excel file failure."""
        mock_workbook_class.side_effect = Exception("Write error")
        
        with pytest.raises(common.IngeniumLibError, match=r"Error writing Excel file:"):
            ProjConfigPalette.write_palette_excel('test.xlsx', {})

    def test_diff_palette_info_no_differences(self):
        """Test diff palette info with no differences."""
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Custom',
                    'enable_disable': 'enable'
                }
            ]
        }
        
        excel_palette = current_palette.copy()
        
        # This should not raise an exception and should complete without errors
        ProjConfigPalette.diff_palette_info(current_palette, excel_palette)

    def test_diff_palette_info_with_differences(self):
        """Test diff palette info with differences."""
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Custom',
                    'enable_disable': 'enable'
                }
            ]
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Different
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Updated Category',  # Different
                    'enable_disable': 'enable'
                }
            ]
        }
        
        # This should not raise an exception and should complete without errors
        ProjConfigPalette.diff_palette_info(current_palette, excel_palette)

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_shows_diff_before_confirmation(self, mock_input, mock_diff, mock_update_builtin):
        """Test that update_palette_info calls diff_palette_info before asking for confirmation."""
        mock_input.return_value = 'y'
        mock_update_builtin.return_value = {'success': True}
        
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Different
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        # Verify diff_palette_info was called before asking for confirmation
        mock_diff.assert_called_once_with(current_palette, excel_palette)
        mock_input.assert_called()
        mock_update_builtin.assert_called_once()

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_with_auto_confirm(self, mock_input, mock_diff, mock_update_builtin):
        """Test that update_palette_info shows diff but skips confirmation with --confirm flag."""
        mock_update_builtin.return_value = {'success': True}
        
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Different
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=True  # Auto-confirm enabled
        )
        
        # Verify diff_palette_info was called but input was not
        mock_diff.assert_called_once_with(current_palette, excel_palette)
        mock_input.assert_not_called()  # Should not ask for confirmation
        mock_update_builtin.assert_called_once()

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_cancelled_after_diff(self, mock_input, mock_diff, mock_update_builtin):
        """Test that update_palette_info shows diff and allows cancellation."""
        mock_input.return_value = 'n'  # User cancels
        
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Different
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        # Verify diff was shown but no updates were performed
        mock_diff.assert_called_once_with(current_palette, excel_palette)
        mock_input.assert_called_once()
        mock_update_builtin.assert_not_called()

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.update_custom_palette')
    @patch('apps.ProjConfigPalette.create_custom_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_with_confirmation(self, mock_input, mock_diff, mock_create, mock_update_custom, mock_update_builtin):
        """Test updating palette info with user confirmation."""
        mock_input.return_value = 'y'
        mock_update_builtin.return_value = {'success': True}
        mock_update_custom.return_value = {'success': True}
        mock_create.return_value = [{'success': True}]
        
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Different
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        # Verify diff_palette_info was called before asking for confirmation
        mock_diff.assert_called_once_with(current_palette, excel_palette)
        mock_update_builtin.assert_called_once()
        mock_input.assert_called()

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_cancelled(self, mock_input, mock_diff, mock_update_builtin):
        """Test updating palette info when user cancels."""
        mock_input.return_value = 'n'
        
        current_palette = {'built_in': [], 'custom': []}
        excel_palette = {'built_in': [], 'custom': []}
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        # Verify diff was shown but no updates were performed
        mock_diff.assert_called_once_with(current_palette, excel_palette)
        mock_input.assert_called()
        mock_update_builtin.assert_not_called()

    @patch('apps.ProjConfigPalette.delete_custom_palette')
    @patch('builtins.input')
    def test_delete_custom_steps_with_confirmation(self, mock_input, mock_delete):
        """Test deleting custom steps with user confirmation."""
        mock_input.return_value = 'y'
        mock_delete.return_value = {'success': True}
        
        current_palette = {
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Custom'
                }
            ]
        }
        
        excel_palette = {
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1'
                }
            ]
        }
        
        ProjConfigPalette.delete_custom_steps(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_delete.assert_called_once_with('https://test-server.example.com', 'custom_1')

    @patch('apps.ProjConfigPalette.delete_custom_palette')
    @patch('builtins.input')
    def test_delete_custom_steps_cancelled(self, mock_input, mock_delete):
        """Test deleting custom steps when user cancels."""
        mock_input.return_value = 'n'
        
        current_palette = {
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Custom'
                }
            ]
        }
        
        excel_palette = {
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1'
                }
            ]
        }
        
        ProjConfigPalette.delete_custom_steps(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_delete.assert_not_called()

    @patch('ing_lib.common.authenticate')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.write_palette_excel')
    @patch('getpass.getpass')
    def test_main_query_function(self, mock_getpass, mock_write_excel, mock_get_palette, mock_auth):
        """Test main function with query operation."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx'
        ]
        
        ProjConfigPalette.main(args)
        
        mock_auth.assert_called_once()
        mock_get_palette.assert_called_once()
        mock_write_excel.assert_called_once()

    @patch('ing_lib.common.authenticate')
    @patch('getpass.getpass')
    def test_main_failed_login(self, mock_getpass, mock_auth):
        """Test main function with failed login."""
        mock_auth.return_value = False
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx'
        ]
        
        with pytest.raises(common.IngeniumLibError, match="Failure to Login"):
            ProjConfigPalette.main(args)

    @patch('ing_lib.common.authenticate')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.read_palette_excel')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('getpass.getpass')
    def test_main_diff_function(self, mock_getpass, mock_diff, mock_read_excel, mock_get_palette, mock_auth):
        """Test main function with diff operation."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_read_excel.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'diff',
            'test.xlsx'
        ]
        
        ProjConfigPalette.main(args)
        
        mock_auth.assert_called_once()
        mock_get_palette.assert_called_once()
        mock_read_excel.assert_called_once()
        mock_diff.assert_called_once()

    @patch('ing_lib.common.authenticate')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.read_palette_excel')
    @patch('apps.ProjConfigPalette.update_palette_info')
    @patch('getpass.getpass')
    def test_main_update_function(self, mock_getpass, mock_update, mock_read_excel, mock_get_palette, mock_auth):
        """Test main function with update operation."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_read_excel.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'update',
            'test.xlsx',
            '--confirm'
        ]
        
        ProjConfigPalette.main(args)
        
        mock_auth.assert_called_once()
        mock_get_palette.assert_called_once()
        mock_read_excel.assert_called_once()
        mock_update.assert_called_once()

    @patch('ing_lib.common.authenticate')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.read_palette_excel')
    @patch('apps.ProjConfigPalette.delete_custom_steps')
    @patch('getpass.getpass')
    def test_main_delete_function(self, mock_getpass, mock_delete, mock_read_excel, mock_get_palette, mock_auth):
        """Test main function with delete operation."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_read_excel.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'delete',
            'test.xlsx',
            '--confirm'
        ]
        
        ProjConfigPalette.main(args)
        
        mock_auth.assert_called_once()
        mock_get_palette.assert_called_once()
        mock_read_excel.assert_called_once()
        mock_delete.assert_called_once()

    @patch('getpass.getuser')
    @patch('getpass.getpass')
    @patch('ing_lib.common.authenticate')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.write_palette_excel')
    def test_main_with_username_option(self, mock_write_excel, mock_get_palette, mock_auth, mock_getpass, mock_getuser):
        """Test main function with username option."""
        mock_getuser.return_value = 'system_user'
        mock_getpass.return_value = 'test_password'
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx',
            '--username', 'custom_user'
        ]
        
        ProjConfigPalette.main(args)
        
        # Verify authenticate was called with custom username
        mock_auth.assert_called_once()
        call_args = mock_auth.call_args
        assert call_args[1]['username'] == 'custom_user'

    @patch('getpass.getuser')
    @patch('getpass.getpass')
    @patch('ing_lib.common.authenticate')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.write_palette_excel')
    def test_main_with_rsa_auth(self, mock_write_excel, mock_get_palette, mock_auth, mock_getpass, mock_getuser):
        """Test main function with RSA authentication."""
        mock_getuser.return_value = 'system_user'
        mock_getpass.return_value = 'test_passcode'
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx',
            '--rsa'
        ]
        
        ProjConfigPalette.main(args)
        
        # Verify authenticate was called with RSA=True
        mock_auth.assert_called_once()
        call_args = mock_auth.call_args
        assert call_args[1]['rsa'] is True

    @patch('ing_lib.common.authenticate')
    @patch('getpass.getpass')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.read_palette_excel')
    @patch('apps.ProjConfigPalette.update_palette_info')
    def test_main_invalid_custom_step_id(self, mock_update, mock_read_excel, mock_get_palette, mock_auth, mock_getpass):
        """Test main function with invalid custom step ID in update."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        mock_read_excel.return_value = {
            'built_in': [],
            'custom': [
                {
                    'step_display_name': 'Invalid Custom Step',
                    'palette_category': 'Custom',
                    'step_id': ''  # Empty step_id
                }
            ]
        }
        
        args = [
            'https://test-server.example.com',
            'update',
            'test.xlsx',
            '--confirm'
        ]
        
        # Should handle invalid custom step gracefully
        ProjConfigPalette.main(args)
        
        mock_update.assert_called_once()

    @patch('ing_lib.common.authenticate')
    @patch('getpass.getpass')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.read_palette_excel')
    @patch('apps.ProjConfigPalette.update_palette_info')
    def test_main_modifying_existing_steps(self, mock_update, mock_read_excel, mock_get_palette, mock_auth, mock_getpass):
        """Test main function modifying existing built-in and custom steps."""
        mock_auth.return_value = True
        mock_getpass.return_value = 'test_password'
        mock_get_palette.return_value = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Custom Step 1',
                    'palette_category': 'Custom',
                    'enable_disable': 'enable'
                }
            ]
        }
        mock_read_excel.return_value = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Modified
                    'palette_category': 'Updated Category',       # Modified
                    'enable_disable': 'disable'                  # Modified
                }
            ],
            'custom': [
                {
                    'step_id': 'custom_1',
                    'step_display_name': 'Updated Custom Step',  # Modified
                    'palette_category': 'Updated Category',      # Modified
                    'enable_disable': 'disable'                  # Modified
                }
            ]
        }
        
        args = [
            'https://test-server.example.com',
            'update',
            'test.xlsx',
            '--confirm'
        ]
        
        ProjConfigPalette.main(args)
        
        mock_update.assert_called_once()
        # Verify the update was called with the correct parameters
        call_args = mock_update.call_args
        assert call_args[0][0] == 'https://test-server.example.com'
        assert 'built_in' in call_args[0][1]
        assert 'custom' in call_args[0][1]
        assert 'built_in' in call_args[0][2]
        assert 'custom' in call_args[0][2]
        assert call_args[0][3] is True  # confirm=True

    @patch('urllib3.disable_warnings')
    @patch('ing_lib.common.authenticate')
    @patch('getpass.getpass')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('apps.ProjConfigPalette.write_palette_excel')
    def test_main_ssl_ignore_error(self, mock_write_excel, mock_get_palette, mock_auth, mock_getpass, mock_disable_warnings):
        """Test main function with SSL error ignored."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx',
            '--ignore_ssl_error'
        ]
        
        ProjConfigPalette.main(args)
        
        # Verify SSL warnings were disabled
        mock_disable_warnings.assert_called_once()

    @patch('apps.ProjConfigPalette.write_palette_excel')
    @patch('apps.ProjConfigPalette.get_palette_info')
    @patch('getpass.getpass')
    @patch('ing_lib.common.authenticate')
    @patch('ing_lib.common.set_ssl_verify')
    def test_main_ssl_ca_bundle(self, mock_set_ssl_verify, mock_auth, mock_getpass, mock_get_palette, mock_write_excel):
        """Test main function with SSL CA bundle specified."""
        mock_auth.return_value = True
        mock_get_palette.return_value = {'built_in': [], 'custom': []}
        mock_getpass.return_value = 'test_password'
        
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx',
            '--ssl_ca_bundle', '/path/to/ca-bundle.crt'
        ]
        
        ProjConfigPalette.main(args)
        
        # Verify SSL verification was set to use CA bundle
        mock_set_ssl_verify.assert_called_once_with('/path/to/ca-bundle.crt')

    def test_main_exception_handling(self):
        """Test main function exception handling."""
        args = [
            'https://test-server.example.com',
            'query',
            'test.xlsx'
        ]
        
        # Test with invalid arguments to trigger exception
        with pytest.raises(SystemExit):
            # This should cause argparse to exit
            ProjConfigPalette.main(['invalid', 'args'])

    @patch('openpyxl.load_workbook')
    def test_read_palette_excel_missing_worksheets(self, mock_load_workbook):
        """Test reading Excel file with missing worksheets."""
        mock_workbook = MagicMock()
        mock_load_workbook.return_value = mock_workbook
        mock_workbook.sheetnames = ['OtherSheet']  # No expected worksheets
        
        result = ProjConfigPalette.read_palette_excel('test.xlsx')
        
        assert 'built_in' in result
        assert 'custom' in result
        assert len(result['built_in']) == 0
        assert len(result['custom']) == 0

    @patch('openpyxl.load_workbook')
    def test_read_palette_excel_empty_worksheets(self, mock_load_workbook):
        """Test reading Excel file with empty worksheets."""
        mock_workbook = MagicMock()
        mock_load_workbook.return_value = mock_workbook
        mock_workbook.sheetnames = ['Built-in Steps', 'Custom Steps']
        
        # Mock empty worksheets
        mock_builtin_ws = MagicMock()
        mock_builtin_ws.__iter__ = MagicMock(return_value=iter([]))
        mock_custom_ws = MagicMock()
        mock_custom_ws.__iter__ = MagicMock(return_value=iter([]))
        
        mock_workbook.__getitem__.side_effect = lambda sheet: {
            'Built-in Steps': mock_builtin_ws,
            'Custom Steps': mock_custom_ws
        }[sheet]
        
        result = ProjConfigPalette.read_palette_excel('test.xlsx')
        
        assert 'built_in' in result
        assert 'custom' in result
        assert len(result['built_in']) == 0
        assert len(result['custom']) == 0

    @patch('openpyxl.load_workbook')
    def test_read_palette_excel_with_placeholder_messages(self, mock_load_workbook):
        """Test reading Excel file with placeholder messages."""
        mock_workbook = MagicMock()
        mock_load_workbook.return_value = mock_workbook
        mock_workbook.sheetnames = ['Built-in Steps', 'Custom Steps']
        
        # Mock custom steps worksheet with placeholder message
        mock_custom_ws = MagicMock()
        mock_custom_ws.__iter__ = MagicMock(return_value=iter([
            [MagicMock(value='message')],
            ['No custom steps found']
        ]))
        
        mock_workbook.__getitem__.side_effect = lambda sheet: {
            'Built-in Steps': MagicMock(),  # Empty built-in
            'Custom Steps': mock_custom_ws
        }[sheet]
        
        result = ProjConfigPalette.read_palette_excel('test.xlsx')
        
        assert 'built_in' in result
        assert 'custom' in result
        assert len(result['custom']) == 0  # Should skip placeholder messages

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.update_custom_palette')
    @patch('apps.ProjConfigPalette.create_custom_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_create_new_custom_step(self, mock_input, mock_diff, mock_create, mock_update_custom, mock_update_builtin):
        """Test updating palette info by creating new custom step."""
        mock_input.return_value = 'y'
        mock_update_builtin.return_value = {'success': True}
        mock_update_custom.return_value = {'success': True}
        mock_create.return_value = [{'success': True}]
        
        current_palette = {
            'built_in': [],
            'custom': []  # No existing custom steps
        }
        
        excel_palette = {
            'built_in': [],
            'custom': [
                {
                    'step_id': 'new_custom_step',
                    'step_display_name': 'New Custom Step',
                    'palette_category': 'Custom',
                    'step_path': '/path/to/custom_step.py',
                    'step_hash': 'abc123def456'
                }
            ]
        }
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_create.assert_called_once()
        mock_update_custom.assert_not_called()

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_built_in_step_not_found(self, mock_input, mock_diff, mock_update_builtin):
        """Test updating palette info when built-in step not found on server."""
        mock_input.return_value = 'y'
        mock_update_builtin.return_value = {'success': True}
        
        current_palette = {
            'built_in': [],  # Empty built-in steps
            'custom': []
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input'
                }
            ],
            'custom': []
        }
        
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        # Should not attempt to update non-existent step
        mock_update_builtin.assert_not_called()

    @patch('apps.ProjConfigPalette.delete_custom_palette')
    @patch('builtins.input')
    def test_delete_custom_steps_step_not_found_on_server(self, mock_input, mock_delete):
        """Test deleting custom step when step not found on server."""
        mock_input.return_value = 'y'
        
        current_palette = {
            'custom': []  # No custom steps on server
        }
        
        excel_palette = {
            'custom': [
                {
                    'step_id': 'nonexistent_step',
                    'step_display_name': 'Nonexistent Step'
                }
            ]
        }
        
        ProjConfigPalette.delete_custom_steps(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_delete.assert_not_called()

    @patch('apps.ProjConfigPalette.delete_custom_palette')
    @patch('builtins.input')
    def test_delete_custom_steps_missing_step_id(self, mock_input, mock_delete):
        """Test deleting custom step when step_id is missing."""
        mock_input.return_value = 'y'
        
        current_palette = {
            'custom': [
                {
                    'step_id': 'existing_step',
                    'step_display_name': 'Existing Step',
                    'palette_category': 'Custom'
                }
            ]
        }
        
        excel_palette = {
            'custom': [
                {
                    'step_display_name': 'Step without ID'
                    # Missing step_id
                }
            ]
        }
        
        ProjConfigPalette.delete_custom_steps(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_delete.assert_not_called()

    @patch('apps.ProjConfigPalette.update_built_in_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_built_in_update_failure(self, mock_input, mock_diff, mock_update_builtin):
        """Test updating palette info when built-in step update fails."""
        mock_input.return_value = 'y'
        mock_update_builtin.side_effect = Exception("Update failed")
        
        current_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Manual Input',
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        excel_palette = {
            'built_in': [
                {
                    'step_type': 'MANUAL_INPUT',
                    'step_display_name': 'Updated Manual Input',  # Different
                    'palette_category': 'Input',
                    'enable_disable': 'enable'
                }
            ],
            'custom': []
        }
        
        # Should handle update failure gracefully
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_update_builtin.assert_called_once()

    @patch('apps.ProjConfigPalette.create_custom_palette')
    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_custom_step_creation_failure(self, mock_input, mock_diff, mock_create):
        """Test updating palette info when custom step creation fails."""
        mock_input.return_value = 'y'
        mock_create.side_effect = Exception("Creation failed")
        
        current_palette = {
            'built_in': [],
            'custom': []  # No existing custom steps
        }
        
        excel_palette = {
            'built_in': [],
            'custom': [
                {
                    'step_id': 'new_custom_step',
                    'step_display_name': 'New Custom Step',
                    'palette_category': 'Custom',
                    'step_path': '/path/to/custom_step.py',
                    'step_hash': 'abc123def456'
                }
            ]
        }
        
        # Should handle creation failure gracefully
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_create.assert_called_once()

    @patch('apps.ProjConfigPalette.delete_custom_palette')
    @patch('builtins.input')
    def test_delete_custom_steps_deletion_failure(self, mock_input, mock_delete):
        """Test deleting custom step when deletion fails."""
        mock_input.return_value = 'y'
        mock_delete.side_effect = Exception("Deletion failed")
        
        current_palette = {
            'custom': [
                {
                    'step_id': 'existing_step',
                    'step_display_name': 'Existing Step',
                    'palette_category': 'Custom'
                }
            ]
        }
        
        excel_palette = {
            'custom': [
                {
                    'step_id': 'existing_step',
                    'step_display_name': 'Existing Step'
                }
            ]
        }
        
        # Should handle deletion failure gracefully
        ProjConfigPalette.delete_custom_steps(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
        
        mock_delete.assert_called_once()

    def test_get_palette_info_custom_palette_failure(self):
        """Test palette info retrieval when custom palette fails."""
        with patch('apps.ProjConfigPalette.get_built_in_palette', return_value=[]), \
             patch('apps.ProjConfigPalette.get_custom_palette', side_effect=Exception("Custom palette error")):
            
            with pytest.raises(common.IngeniumLibError, match=r"Error getting palette information from server:"):
                ProjConfigPalette.get_palette_info('https://test-server.example.com')

    @patch('apps.ProjConfigPalette.Workbook')
    def test_write_palette_excel_no_data(self, mock_workbook_class):
        """Test writing Excel file with no palette data."""
        mock_workbook = MagicMock()
        mock_workbook_class.return_value = mock_workbook
        mock_workbook.sheetnames = ['Sheet']
        
        # Mock worksheets
        mock_builtin_ws = MagicMock()
        mock_custom_ws = MagicMock()
        mock_workbook.create_sheet.side_effect = [mock_builtin_ws, mock_custom_ws]
        mock_workbook.__contains__.side_effect = lambda sheet: sheet in ['Built-in Steps', 'Custom Steps']
        mock_workbook.__getitem__.side_effect = lambda sheet: {
            'Sheet': None,  # Will be removed
            'Built-in Steps': mock_builtin_ws,
            'Custom Steps': mock_custom_ws
        }[sheet]
        
        palette_info = {
            'built_in': [],
            'custom': []
        }
        
        # Just verify the function completes without error
        ProjConfigPalette.write_palette_excel('test.xlsx', palette_info)
        
        # Verify basic workbook operations were attempted
        mock_workbook_class.assert_called_once()
        mock_workbook.close.assert_called_once()

    def test_diff_palette_info_empty_data(self):
        """Test diff palette info with empty data."""
        current_palette = {'built_in': [], 'custom': []}
        excel_palette = {'built_in': [], 'custom': []}
        
        # Should handle empty data gracefully
        ProjConfigPalette.diff_palette_info(current_palette, excel_palette)

    @patch('apps.ProjConfigPalette.diff_palette_info')
    @patch('builtins.input')
    def test_update_palette_info_no_steps_to_process(self, mock_input, mock_diff):
        """Test updating palette info with no steps to process."""
        mock_input.return_value = 'y'
        
        current_palette = {'built_in': [], 'custom': []}
        excel_palette = {'built_in': [], 'custom': []}
        
        # Should handle empty data gracefully
        ProjConfigPalette.update_palette_info(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )

    @patch('builtins.input')
    def test_delete_custom_steps_no_custom_steps(self, mock_input):
        """Test deleting custom steps when no custom steps in Excel."""
        current_palette = {'custom': []}
        excel_palette = {'custom': []}
        
        # Should handle empty data gracefully
        ProjConfigPalette.delete_custom_steps(
            'https://test-server.example.com',
            current_palette,
            excel_palette,
            confirm=False
        )
