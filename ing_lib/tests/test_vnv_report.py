"""
Tests for VnVReport.py app
"""

import pytest
from openpyxl import load_workbook

# Import the module under test
import ing_lib.common as common
from apps.VnVReport import get_input, build_report_data, output_vnv_report


def make_step(step_type, status='PASS', results=None, user_input=None,
              procedure_id='test-proc-1', procedure_title='Test Proc One', number='1'):
    """Helper to build an as-run element dict shaped like the executions API response."""
    return {
        'elem_type': 'STEP',
        'step_type': step_type,
        'procedure_id': procedure_id,
        'procedure_title': procedure_title,
        'number': number,
        'execution_user_input': user_input or {},
        'execution': {
            'meta_data': {'status': status},
            'results': results or {}
        }
    }


SAMPLE_EXECUTION_INFO = {
    'execution_id': 'test-ingenium-10001',
    'time_started': '2025-06-01T10:00:00.000Z',
    'time_completed': '2025-06-01T15:30:00.000Z',
    'venue_name': 'Test Venue',
    'description': 'Test execution description',
    'test_conductors': ['tc1', 'tc2'],
    'used_procedures': [
        {
            'procedure_id': 'test-proc-1',
            'version': 2,
            'run_for_score': True,
            'institutional_id': 'P-1001',
            'institutional_release_id': 'RevA'
        },
        {
            'procedure_id': 'test-proc-2',
            'version': 1,
            'run_for_score': False,
            'institutional_id': 'P-1002',
            'institutional_release_id': ''
        }
    ]
}

SAMPLE_ELEMENTS = [
    # Non-step element - should be skipped
    {'elem_type': 'SECTION'},
    # GDS_MANUAL - collect unique session ids (skip '-1', missing, duplicates)
    make_step('GDS_MANUAL', user_input={'entries': [
        {'session_id': '1234'},
        {'session_id': '-1'},
        {'session_id': '1234'},
        {'name': 'no_session'},
        {'session_id': '5678'},
    ]}),
    # GDS_MANUAL that did not run - should not collect its session
    make_step('GDS_MANUAL', status='NONE', user_input={'entries': [{'session_id': '9999'}]}),
    # MANUAL_INPUT - collect the activity report number
    make_step('MANUAL_INPUT', results={'entries': [
        {'name': 'Activity Report Number', 'actual_value': 'AR-42'},
        {'name': 'Notes', 'actual_value': 'not an ar'},
    ]}),
    # VENUE_CONFIG_MANUAL - collect venue config details and the ampcs machine
    make_step('VENUE_CONFIG_MANUAL', results={
        'fsw_dictionary': 'flight-v1',
        'fsw_version': 'fsw-2.0',
        'gds_version': 'gds-9',
        'sse_dictionary': 'sse-v1',
        'sse_version': 'sse-3',
        'venue_ampcs_address': 'https://ampcs-host:8443/ampcs',
    }),
    # VI status steps - matched and unmatched to used_procedures
    make_step('VERIFICATION_ITEM_STATUS', number='10', results={'vis': [
        {'vi_type': 'TEST', 'vi_id': 'VI-1', 'vi_name': 'First VI'},
        {'vi_type': 'REQ', 'vi_id': 'VI-2', 'vi_name': 'Second VI'},
    ]}),
    make_step('VERIFICATION_ITEM_STATUS', number='4', status='FAIL',
              procedure_id='unknown-proc', procedure_title='Mystery Proc',
              results={'vis': [{'vi_type': 'TEST', 'vi_id': 'VI-3', 'vi_name': 'Third VI'}]}),
]


class TestVnVReportInput:
    """Test get_input for VnVReport."""

    def test_get_input_required_args(self):
        args = [
            'https://test-server.example.com',
            'test-ingenium-10001',
            '/path/to/report.xlsx'
        ]
        inputs = get_input(args)

        assert inputs.ingenium_server == 'https://test-server.example.com'
        assert inputs.execution_id == 'test-ingenium-10001'
        assert inputs.output_file == '/path/to/report.xlsx'
        assert inputs.level == 'SYS'
        assert inputs.ignore_ssl_error is False
        assert inputs.rsa is False

    def test_get_input_optional_args(self):
        args = [
            'https://test-server.example.com',
            'test-ingenium-10001',
            '/path/to/report.xlsx',
            '--username', 'test_user', '--level', 'AVS',
            '--ignore_ssl_error', '--rsa', '--ssl_ca_bundle', '/path/to/ca.pem'
        ]
        inputs = get_input(args)

        assert inputs.username == 'test_user'
        assert inputs.level == 'AVS'
        assert inputs.ignore_ssl_error is True
        assert inputs.rsa is True
        assert inputs.ssl_ca_bundle == '/path/to/ca.pem'


class TestBuildReportData:
    """Test build_report_data assembly of VI/session/config info."""

    def test_build_report_data_vis(self):
        report = build_report_data(SAMPLE_ELEMENTS, SAMPLE_EXECUTION_INFO)

        assert len(report['vis']) == 3

        vi1 = report['vis'][0]
        assert vi1['vi_type'] == 'TEST'
        assert vi1['vi_id'] == 'VI-1'
        assert vi1['vi_name'] == 'First VI'
        assert vi1['status'] == 'PASS'
        assert vi1['step_number'] == '10'
        assert vi1['procedure'] == 'test-proc-1 Test Proc One'
        # Matched procedure gets institutional id + release id
        assert vi1['inst_procedure'] == 'P-1001 RevA'

        # Unmatched procedure gets 'TBD'
        vi3 = report['vis'][2]
        assert vi3['inst_procedure'] == 'TBD'
        assert vi3['status'] == 'FAIL'

    def test_build_report_data_no_release_id(self):
        """A matched procedure with empty release id uses just the institutional id."""
        elements = [make_step('VERIFICATION_ITEM_STATUS', procedure_id='test-proc-2',
                              procedure_title='Test Proc Two', number='5',
                              results={'vis': [{'vi_type': 'TEST', 'vi_id': 'VI-9', 'vi_name': 'VI Nine'}]})]
        report = build_report_data(elements, SAMPLE_EXECUTION_INFO)

        assert report['vis'][0]['inst_procedure'] == 'P-1002'

    def test_build_report_data_sessions_and_config(self):
        report = build_report_data(SAMPLE_ELEMENTS, SAMPLE_EXECUTION_INFO)

        # Only executed GDS_MANUAL sessions, de-duplicated, without '-1'/missing
        assert report['gds_sessions'] == ['1234', '5678']
        assert report['activity_reports'] == ['AR-42']
        assert report['test_venue_addresses'] == ['ampcs-host']
        assert report['fsw_dictionaries'] == ['flight-v1']
        assert report['fsw_versions'] == ['fsw-2.0']
        assert report['gds_versions'] == ['gds-9']
        assert report['sse_dictionaries'] == ['sse-v1']
        assert report['sse_versions'] == ['sse-3']

    def test_build_report_data_preserves_execution_info(self):
        report = build_report_data(SAMPLE_ELEMENTS, SAMPLE_EXECUTION_INFO)

        assert report['execution_id'] == 'test-ingenium-10001'
        assert report['venue_name'] == 'Test Venue'
        assert report['used_procedures'] == SAMPLE_EXECUTION_INFO['used_procedures']
        # The input dict is deep-copied, not mutated
        assert 'vis' not in SAMPLE_EXECUTION_INFO

    def test_build_report_data_empty(self):
        report = build_report_data([], {'execution_id': 'e-1', 'used_procedures': []})

        assert report['vis'] == []
        assert report['gds_sessions'] == []
        assert report['activity_reports'] == []
        assert report['test_venue_addresses'] == []

    def test_build_report_data_no_sessions_or_addresses(self):
        """Steps without GDS sessions or venue addresses produce empty collections."""
        elements = [
            make_step('MANUAL_VERIFICATION', results={}),
            make_step('VENUE_CONFIG_MANUAL', results={
                'fsw_version': 'fsw-2.0',
                # no venue_ampcs_address in results
            }),
            make_step('MANUAL_INPUT', results={'entries': [
                {'name': 'Notes', 'actual_value': 'not an ar'},
            ]}),
        ]
        report = build_report_data(elements, SAMPLE_EXECUTION_INFO)

        assert report['gds_sessions'] == []
        assert report['test_venue_addresses'] == []
        assert report['activity_reports'] == []
        assert report['fsw_versions'] == ['fsw-2.0']


class TestOutputVnVReport:
    """Test output_vnv_report XLSX generation."""

    def build_report_info(self):
        return build_report_data(SAMPLE_ELEMENTS, SAMPLE_EXECUTION_INFO)

    def test_verification_report_sheet(self, tmp_path):
        out_file = tmp_path / "vnv_report.xlsx"
        output_vnv_report(self.build_report_info(), str(out_file), 'SYS')

        wb = load_workbook(str(out_file))
        ws = wb['Verification Report']

        # Execution info block
        assert ws.cell(row=1, column=1).value == 'Ingenium Execution'
        assert ws.cell(row=1, column=2).value == 'test-ingenium-10001'
        assert ws.cell(row=2, column=2).value == '2025-06-01T10:00:00.000Z'
        assert ws.cell(row=4, column=2).value == 'Test Venue'
        assert ws.cell(row=6, column=2).value == 'tc1,tc2'

        # Procedures table
        assert ws.cell(row=7, column=1).value == 'Procedures'
        assert ws.cell(row=8, column=1).value == 'Procedure ID'
        assert ws.cell(row=8, column=4).value == 'P#'
        assert ws.cell(row=9, column=1).value == 'test-proc-1'
        assert ws.cell(row=9, column=2).value == 2
        assert ws.cell(row=9, column=3).value == 'True'
        assert ws.cell(row=9, column=4).value == 'P-1001'
        assert ws.cell(row=9, column=5).value == 'RevA'
        assert ws.cell(row=10, column=1).value == 'test-proc-2'

        # VI results table
        assert ws.cell(row=12, column=1).value == 'VI Type'
        assert ws.cell(row=12, column=6).value == 'Status'
        assert ws.cell(row=13, column=1).value == 'TEST'
        assert ws.cell(row=13, column=2).value == 'VI-1'
        assert ws.cell(row=13, column=4).value == 'test-proc-1 Test Proc One'
        assert ws.cell(row=13, column=5).value == '10'
        assert ws.cell(row=13, column=6).value == 'PASS'
        assert ws.cell(row=15, column=2).value == 'VI-3'

        wb.close()

    def test_vits_sheet(self, tmp_path):
        out_file = tmp_path / "vnv_report.xlsx"
        output_vnv_report(self.build_report_info(), str(out_file), 'SYS')

        wb = load_workbook(str(out_file))
        ws = wb['VITS']

        # Headers
        assert ws.cell(row=1, column=1).value == 'Element Type'
        assert ws.cell(row=1, column=9).value == 'Test Date'
        assert ws.cell(row=1, column=12).value == 'Test Venue Address'
        assert ws.cell(row=1, column=14).value == 'Test Procedure'
        assert ws.cell(row=1, column=24).value == 'Other Issue'

        # First VI row (SYS -> System I&T, test date from time_completed)
        assert ws.cell(row=2, column=1).value == 'TEST'
        assert ws.cell(row=2, column=2).value == 'VI-1'
        assert ws.cell(row=2, column=4).value == 'System I&T'
        assert ws.cell(row=2, column=5).value == 'fsw-2.0'
        assert ws.cell(row=2, column=6).value == 'flight-v1'
        assert ws.cell(row=2, column=7).value == 'sse-3'
        assert ws.cell(row=2, column=8).value == 'gds-9'
        assert ws.cell(row=2, column=9).value == '2025-06-01'
        assert ws.cell(row=2, column=10).value == 'PASS'
        assert ws.cell(row=2, column=11).value == 'AR-42'
        assert ws.cell(row=2, column=12).value == 'ampcs-host'
        assert ws.cell(row=2, column=13).value == '1234,5678'
        assert ws.cell(row=2, column=14).value == 'P-1001 RevA'

        # Second VI row
        assert ws.cell(row=3, column=2).value == 'VI-2'

        # Unmatched procedure -> TBD
        assert ws.cell(row=4, column=14).value == 'TBD'

        wb.close()

    def test_vits_sheet_levels(self, tmp_path):
        report = self.build_report_info()

        for level, expected in [('AVS', 'Avionics System I&T'), ('FSW', 'FSW I&T')]:
            out_file = tmp_path / f"vnv_report_{level}.xlsx"
            output_vnv_report(report, str(out_file), level)

            wb = load_workbook(str(out_file))
            assert wb['VITS'].cell(row=2, column=4).value == expected
            wb.close()

    def test_vits_sheet_bad_time_completed(self, tmp_path):
        report = self.build_report_info()
        report['time_completed'] = None
        out_file = tmp_path / "vnv_report.xlsx"

        output_vnv_report(report, str(out_file), 'SYS')

        wb = load_workbook(str(out_file))
        assert wb['VITS'].cell(row=2, column=9).value == 'TBD'
        wb.close()

    def test_output_empty_vis(self, tmp_path):
        report = build_report_data([], SAMPLE_EXECUTION_INFO)
        out_file = tmp_path / "empty_report.xlsx"

        output_vnv_report(report, str(out_file), 'SYS')

        wb = load_workbook(str(out_file))
        ws = wb['Verification Report']
        assert ws.cell(row=12, column=1).value == 'VI Type'
        assert ws.max_row == 12

        vits = wb['VITS']
        assert vits.cell(row=1, column=1).value == 'Element Type'
        assert vits.max_row == 1

        wb.close()

    def test_vits_sheet_no_sessions_or_addresses(self, tmp_path):
        """VITS cells are empty when no GDS sessions or venue addresses exist."""
        elements = [make_step('VERIFICATION_ITEM_STATUS', number='10', results={'vis': [
            {'vi_type': 'TEST', 'vi_id': 'VI-1', 'vi_name': 'First VI'},
        ]})]
        report = build_report_data(elements, SAMPLE_EXECUTION_INFO)
        out_file = tmp_path / "vnv_report.xlsx"

        output_vnv_report(report, str(out_file), 'SYS')

        wb = load_workbook(str(out_file))
        ws = wb['VITS']
        assert ws.cell(row=2, column=12).value in ('', None)
        assert ws.cell(row=2, column=13).value in ('', None)
        wb.close()

    def test_output_missing_collection_keys(self, tmp_path):
        """Report writes without error when collection keys are absent."""
        report_info = {
            'execution_id': 'test-ingenium-10001',
            'time_started': '2025-06-01T10:00:00.000Z',
            'time_completed': '2025-06-01T15:30:00.000Z',
            'venue_name': 'Test Venue',
            'description': 'Test execution description',
            'test_conductors': ['tc1'],
            'used_procedures': [],
            'vis': [{'vi_type': 'TEST', 'vi_id': 'VI-1', 'vi_name': 'First VI',
                     'status': 'PASS', 'procedure': 'test-proc-1 Test Proc One',
                     'ing_procedure': 'test-proc-1', 'inst_procedure': 'TBD',
                     'step_number': '10'}]
        }
        out_file = tmp_path / "vnv_report.xlsx"

        output_vnv_report(report_info, str(out_file), 'SYS')

        wb = load_workbook(str(out_file))
        ws = wb['VITS']
        assert ws.cell(row=2, column=2).value == 'VI-1'
        assert ws.cell(row=2, column=12).value in ('', None)
        assert ws.cell(row=2, column=13).value in ('', None)
        wb.close()

    def test_output_bad_path(self):
        with pytest.raises(common.IngeniumLibError) as exc_info:
            output_vnv_report(self.build_report_info(), '/nonexistent_dir/report.xlsx', 'SYS')

        assert "Unable to save" in str(exc_info.value)
