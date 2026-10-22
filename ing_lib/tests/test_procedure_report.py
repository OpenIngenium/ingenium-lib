"""
Tests for ProcedureReport.py app
"""

import pytest
from openpyxl import load_workbook

# Import the module under test
import ing_lib.common as common
from apps.ProcedureReport import get_input, build_report


class MockInputs:
    """Mock argparse inputs for ProcedureReport."""
    def __init__(self, output_file):
        self.output_file = output_file


SAMPLE_PROC_VERS = [
    {
        'procedure': {
            'procedure_id': 'test-procedure-1',
            'title': 'Test Procedure One',
            'institutional_id': 'INST-001',
            'author': 'testuser',
            'time_created': '2025-01-15T10:30:00.000Z',
            'status': 'ACTIVE'
        },
        'version': {
            'version': 1,
            'institutional_release_id': 'REL-1',
            'time_saved': '2025-02-01T08:00:00.000Z',
            'status': 'RELEASED',
            'flight_dictionary_version': 'flight-v1.0',
            'sse_dictionary_version': 'sse-v2.0'
        }
    },
    {
        'procedure': {
            'procedure_id': 'test-procedure-2',
            'title': 'Test Procedure Two',
            'institutional_id': 'INST-002',
            'author': 'testuser',
            'time_created': '2025-03-10T12:00:00.000Z',
            'status': 'ACTIVE'
        },
        'version': {
            'version': 3,
            'institutional_release_id': 'REL-2',
            'time_saved': '2025-04-01T09:30:00.000Z',
            'status': 'RELEASED',
            'flight_dictionary_version': 'flight-v1.0',
            'sse_dictionary_version': 'sse-v2.1'
        }
    }
]


class TestProcedureReportInput:
    """Test get_input for ProcedureReport."""

    def test_get_input_required_args(self):
        args = [
            'https://test-server.example.com',
            '/path/to/report.xlsx'
        ]
        inputs = get_input(args)

        assert inputs.ingenium_server == 'https://test-server.example.com'
        assert inputs.output_file == '/path/to/report.xlsx'
        assert inputs.debug is False
        assert inputs.ignore_ssl_error is False
        assert inputs.rsa is False

    def test_get_input_optional_args(self):
        args = [
            'https://test-server.example.com',
            '/path/to/report.xlsx',
            '--debug', '--username', 'test_user',
            '--ignore_ssl_error', '--rsa'
        ]
        inputs = get_input(args)

        assert inputs.debug is True
        assert inputs.username == 'test_user'
        assert inputs.ignore_ssl_error is True
        assert inputs.rsa is True


class TestBuildReport:
    """Test build_report XLSX generation."""

    def test_build_report_contents(self, tmp_path):
        """Test the report writes headers and procedure/version rows correctly."""
        out_file = tmp_path / "report.xlsx"
        inputs = MockInputs(str(out_file))

        build_report(SAMPLE_PROC_VERS, inputs)

        wb = load_workbook(str(out_file))
        ws = wb['Procedure Version Report']

        # Headers
        assert ws.cell(row=1, column=1).value == 'Procedure ID'
        assert ws.cell(row=1, column=6).value == 'Procedure Status'
        assert ws.cell(row=1, column=12).value == 'SSE Dictionary'

        # Row 2 (first procedure/version)
        assert ws.cell(row=2, column=1).value == 'test-procedure-1'
        assert ws.cell(row=2, column=2).value == 'Test Procedure One'
        assert ws.cell(row=2, column=5).value == '2025-01-15'
        assert ws.cell(row=2, column=7).value == 1
        assert ws.cell(row=2, column=8).value == 'REL-1'
        assert ws.cell(row=2, column=11).value == 'flight-v1.0'

        # Row 3 (second procedure/version)
        assert ws.cell(row=3, column=1).value == 'test-procedure-2'
        assert ws.cell(row=3, column=7).value == 3
        assert ws.cell(row=3, column=10).value == 'RELEASED'

        # Autofilter covers the data range
        assert ws.auto_filter.ref == 'A1:L3'

        wb.close()

    def test_build_report_empty(self, tmp_path):
        """Test the report with no procedures still writes headers."""
        out_file = tmp_path / "empty_report.xlsx"
        inputs = MockInputs(str(out_file))

        build_report([], inputs)

        wb = load_workbook(str(out_file))
        ws = wb['Procedure Version Report']

        assert ws.cell(row=1, column=1).value == 'Procedure ID'
        assert ws.max_row == 1
        assert ws.auto_filter.ref == 'A1:L1'

        wb.close()

    def test_build_report_bad_path(self):
        """Test saving to an invalid path raises IngeniumLibError."""
        inputs = MockInputs('/nonexistent_dir/report.xlsx')

        with pytest.raises(common.IngeniumLibError) as exc_info:
            build_report(SAMPLE_PROC_VERS, inputs)

        assert "Unable to save" in str(exc_info.value)
