"""
This script creates a report of all the procedures/versions managed by an Ingenium Instance

Authors:
    * Chris Swan (christopher.a.swan@jpl.nasa.gov)
"""

##################################################### Imports ########################################################
import logging
from ing_lib.logs import init_console_logger,get_logger
init_console_logger()

import ing_lib.common as common
from ing_lib.procedures import get_procedures_and_versions
import argparse
import getpass
import urllib3
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter


##################################################### Functions ######################################################

logger = get_logger(__name__)

def get_input(args=[]):
    """
    This function gathers inputs for the Ingenium Procedure Report Script

    Parameters
    --------
    args
        list of input arguments. Used when this is called from another Python module.

    Returns
    -------
    inputs: object
        Argparse input object
    """

    parser = argparse.ArgumentParser(
        description="This script creates a report on the Procedures/Versions managed by Ingenium",
        prog="ProcedureReport.py")
    parser.add_argument('ingenium_server', type=str,
                        help='Full Path to Ingenium Server (e.g. https://ingenium-psyche.jpl.nasa.gov')
    parser.add_argument('output_file', type=str, help='Output XLSX File')
    parser.add_argument('--ignore_ssl_error', action='store_true', help='Ignore SSL verification error')
    parser.add_argument('--ssl_ca_bundle', type=str,
                        help='Path to the SSL CA bundle. If provided, this will override ignore_ssl_error.')
    parser.add_argument('--debug', action='store_true', help='Enables debug output (raw JSON)')
    parser.add_argument('--username', type=str,
                        help='Optional input to use a different username to login to Ingenium server.')
    parser.add_argument('--rsa', action='store_true',
                        help='Uses RSA Two Factor Authentication (username/passcode) to authenticate.')

    if len(args) > 0:
        inputs = parser.parse_args(args)
    else:
        inputs = parser.parse_args()

    # Setup debug logging (if desired)
    if inputs.debug:
        get_logger().setLevel(logging.DEBUG)
        for handler in get_logger().handlers:
            handler.setLevel(logging.DEBUG)
            logger.debug("Logging set to Debug.")

    return inputs


def build_report(output, inputs):
    """
    This function builds the XLSX report containing the procedure/version information.

    Parameters
    ----------
    output: list
         List of dictionaries containing the procedure and version information
    inputs: object
        Argparse input object

    Returns
    -------

    """
    # Create a workbook and add a worksheet.
    try:
        workbook = Workbook()
    except:
        msg = f"Failure to open XLSX workbook ({inputs.output_file}). Can not proceed. Exiting."
        logger.error(msg, exc_info=True)
        raise common.IngeniumLibError(msg)

    worksheet = workbook.active
    worksheet.title = 'Procedure Version Report'

    # Add formats
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))
    title_font = Font(color='FF000000', bold=True)
    title_fill = PatternFill(fill_type='solid', start_color='FFBFBFBF', end_color='FFBFBFBF')
    title_alignment = Alignment(wrap_text=True)
    info_font = Font(color='FF000000', bold=False)
    info_fill = PatternFill(fill_type='solid', start_color='FFFFFFFF', end_color='FFFFFFFF')
    info_alignment = Alignment(wrap_text=True, horizontal='center')

    def write_cell(row, col, value, font, fill, alignment):
        cell = worksheet.cell(row=row, column=col, value=value)
        cell.font = font
        cell.fill = fill
        cell.alignment = alignment
        cell.border = thin_border

    # Set Column Widths
    column_widths = [25, 50, 25, 20, 25, 15, 15, 25, 25, 15, 30, 30]
    for idx, width in enumerate(column_widths, 1):
        worksheet.column_dimensions[get_column_letter(idx)].width = width

    row = 1

    # Set Titles
    headers = ['Procedure ID',
               'Procedure Title',
               'Procedure Institutional ID',
               'Procedure Author',
               'Procedure Creation Date',
               'Procedure Status',
               'Version ID',
               'Institutional Release ID',
               'Version Creation Date',
               'Version Status',
               'Flight Dictionary',
               'SSE Dictionary']

    for col_idx, header in enumerate(headers, 1):
        write_cell(row, col_idx, header, title_font, title_fill, title_alignment)

    # write output
    for proc_ver in output:
        row = row + 1
        write_cell(row, 1, proc_ver['procedure']['procedure_id'], info_font, info_fill, info_alignment)
        write_cell(row, 2, proc_ver['procedure']['title'], info_font, info_fill, info_alignment)
        write_cell(row, 3, proc_ver['procedure']['institutional_id'], info_font, info_fill, info_alignment)
        write_cell(row, 4, proc_ver['procedure']['author'], info_font, info_fill, info_alignment)
        write_cell(row, 5, proc_ver['procedure']['time_created'][:10], info_font, info_fill, info_alignment)
        write_cell(row, 6, proc_ver['procedure']['status'], info_font, info_fill, info_alignment)
        write_cell(row, 7, proc_ver['version']['version'], info_font, info_fill, info_alignment)
        write_cell(row, 8, proc_ver['version']['institutional_release_id'], info_font, info_fill, info_alignment)
        write_cell(row, 9, proc_ver['version']['time_saved'][:10], info_font, info_fill, info_alignment)
        write_cell(row, 10, proc_ver['version']['status'], info_font, info_fill, info_alignment)
        write_cell(row, 11, proc_ver['version']['flight_dictionary_version'], info_font, info_fill, info_alignment)
        write_cell(row, 12, proc_ver['version']['sse_dictionary_version'], info_font, info_fill, info_alignment)

    # Add Filter
    worksheet.auto_filter.ref = f"A1:L{row}"

    try:
        workbook.save(inputs.output_file)
        workbook.close()
    except:
        msg = 'Unable to save Procedure / Version Report.'
        logger.error(msg, exc_info=True)
        raise common.IngeniumLibError(msg)

    return

##################################################### Main ###########################################################


def main(args=[]):
    """
    This is the main function of the Ingenium Procedure Report Generator

    Parameters
    ----------
    args
        Array of input arguments. Used when this function is called from another Python module.

    Returns
    -------
        None
    """

    # Gets initial input
    inputs = get_input(args)

    # if ssl_ca_bundle is specified, it will take precedence over ignore_ssl_error
    if inputs.ssl_ca_bundle is not None:
        common.set_ssl_verify(inputs.ssl_ca_bundle)
    else:
        common.set_ssl_verify(not inputs.ignore_ssl_error)
        if not common.get_ssl_verify():
            # To suppress SSL warnings
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    logger.debug(f"Ingenium.ssl_verify:{common.get_ssl_verify()}")

    if inputs.username:
        username = inputs.username
    else:
        username = getpass.getuser()

    if inputs.rsa:
        pw_prompt = f"Enter RSA Passcode for {username}:"
    else:
        pw_prompt = f"Enter LDAP Password for {username}:"

    # Login to Ingenium
    login = common.authenticate(inputs.ingenium_server, username=username,
                                password=getpass.getpass(pw_prompt), force=True,
                                rsa=inputs.rsa)
    if not login:
        msg = 'Failure to Login - Can not proceed. Exiting.'
        logger.error(msg)
        raise common.IngeniumLibError(msg)


    # Query/Assemble the Data
    proc_vers = get_procedures_and_versions(inputs.ingenium_server)

    # Output to XLSX
    build_report(proc_vers, inputs)


if __name__ == '__main__':
    try:
        main()
    except common.IngeniumLibError:
        logger.error("Ingenium Procedure Report encountered an error and needs to exit. Please check the log messages for the source of the error.")
