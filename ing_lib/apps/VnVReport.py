"""
This script creates a report of all the Verification Items addressed in an Ingenium Execution and their status

Authors:
    * Chris Swan (christopher.a.swan@jpl.nasa.gov)
"""

##################################################### Imports ########################################################
from ing_lib.logs import init_console_logger,get_logger
init_console_logger()

import ing_lib.common as common
from ing_lib.executions import get_ingenium_execution_elements
import argparse
import copy
import getpass
import urllib3
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

##################################################### Constants ######################################################

TEST_LEVELS = {'SYS': 'System I&T', 'AVS': 'Avionics System I&T', 'FSW': 'FSW I&T'}

##################################################### Functions ######################################################

logger = get_logger(__name__)

def get_input(args=[]):
    """
    This function gathers inputs for the Ingenium V&V Report Script

    Parameters
    --------
    args
        list of input arguments. Used when this is called from another Python module.

    Returns
    -------
    inputs: object
        Argparse input object
    """

    parser = argparse.ArgumentParser(description="This script creates a report on the V&V results for an Ingenium Execution.",
                                     prog="VnVReport.py")
    parser.add_argument('ingenium_server', type=str,
                        help='Full Path to Ingenium Server (e.g. https://ingenium-psyche.jpl.nasa.gov)')
    parser.add_argument('execution_id', type=str, help='The Ingenium Execution ID (e.g. psyche-ingenium-10311')
    parser.add_argument('output_file', type=str, help='Output XLSX File')
    parser.add_argument('--ignore_ssl_error', action='store_true', help='Ignore SSL verification error')
    parser.add_argument('--ssl_ca_bundle', type=str,
                        help='Path to the SSL CA bundle. If provided, this will override ignore_ssl_error.')
    parser.add_argument('--level', default='SYS', choices=['SYS', 'AVS', 'FSW'],
                        help='What level of testing is being performed (default SYS/System) to use for DVI inputs.')
    parser.add_argument('--username', type=str,
                        help='Optional input to use a different username to login to Ingenium server.')
    parser.add_argument('--rsa', action='store_true',
                        help='Uses RSA Two Factor Authentication (username/passcode) to authenticate.')

    if len(args) > 0:
        inputs = parser.parse_args(args)
    else:
        inputs = parser.parse_args()

    return inputs


def build_report_data(as_run_elements, execution_information):
    """
    This function assembles the step information from the as-run data.

    Parameters
    ----------
    as_run_elements: list
        List of element objects (dictionaries) in the execution
    execution_information: dict
        Dictionary of information regarding the execution

    Returns
    -------
    report_data: dict
        Dictionary of information used to build the report
    """

    report_data = copy.deepcopy(execution_information)

    vis = []
    gds_sessions = []
    test_venue_addresses = []
    activity_reports = []
    fsw_dictionaries = []
    fsw_versions = []
    gds_versions = []
    sse_dictionaries = []
    sse_versions = []

    # Walk through the elements and report any VI Status Results
    for elem in as_run_elements:

        # If the element is an Ingenium Step (not section,etc)
        if elem.get('elem_type') != 'STEP':
            continue

        step_type = elem.get('step_type')
        execution = elem.get('execution') or {}
        meta_data = execution.get('meta_data') or {}
        results = execution.get('results') or {}

        # If the status is NONE the step did not execute (so it can be skipped)
        if meta_data.get('status') != 'NONE':
            # Check if manual GDS Step
            # If the step in a MANUAL GDS step - capture the session IDs and the executions they occurred in
            if step_type == 'GDS_MANUAL':
                for entry in (elem.get('execution_user_input') or {}).get('entries', []):
                    session_id = entry.get('session_id')
                    # If the session_id exists
                    if session_id is not None:
                        session_id = str(session_id)
                        if session_id not in gds_sessions and session_id != '-1':
                            gds_sessions.append(session_id)

            # If the step is a manual input step - check if the any of the entries hold activity report information
            # This relies on the entry name containing "activity report"
            if step_type == 'MANUAL_INPUT':
                for entry in results.get('entries', []):
                    if 'activity report' in str(entry.get('name', '')).lower():
                        ar_num = str(entry.get('actual_value'))
                        if ar_num not in activity_reports:
                            activity_reports.append(ar_num)

            # Now extract configuration information from venue config steps
            if step_type == 'VENUE_CONFIG_MANUAL':
                for result_key, collection in (('fsw_dictionary', fsw_dictionaries),
                                               ('fsw_version', fsw_versions),
                                               ('gds_version', gds_versions),
                                               ('sse_dictionary', sse_dictionaries),
                                               ('sse_version', sse_versions)):
                    value = results.get(result_key)
                    if value and value not in collection:
                        collection.append(value)

            # Check the venue it executed on (if available) and add it to the list.
            if results.get('venue_ampcs_address'):
                venue_address = results['venue_ampcs_address'].split('//')[1].split(':')[0]
                if venue_address not in test_venue_addresses:
                    test_venue_addresses.append(venue_address)

        # Check if the step type = VI Status
        if step_type == 'VERIFICATION_ITEM_STATUS':
            procedure_id = elem.get('procedure_id')

            # The P Number is not available via the VI step so we will look it up
            inst_procedure = 'TBD'
            for procedure in report_data.get('used_procedures', []):
                if procedure_id == procedure.get('procedure_id'):
                    inst_procedure = str(procedure.get('institutional_id'))
                    # if there is a revision - add that too
                    if procedure.get('institutional_release_id'):
                        inst_procedure = inst_procedure + ' ' + str(procedure['institutional_release_id'])
                    break

            for vi in results.get('vis', []):
                verification_status = {'status': meta_data.get('status'),
                                       'procedure': str(procedure_id) + " " + str(elem.get('procedure_title')),
                                       'ing_procedure': procedure_id,
                                       'inst_procedure': inst_procedure,
                                       'step_number': elem.get('number'),
                                       'vi_type': vi.get('vi_type'),
                                       'vi_id': vi.get('vi_id'),
                                       'vi_name': vi.get('vi_name')}
                vis.append(verification_status)

    report_data['vis'] = vis
    report_data['gds_sessions'] = gds_sessions
    report_data['test_venue_addresses'] = test_venue_addresses
    report_data['activity_reports'] = activity_reports
    report_data['fsw_dictionaries'] = fsw_dictionaries
    report_data['fsw_versions'] = fsw_versions
    report_data['gds_versions'] = gds_versions
    report_data['sse_dictionaries'] = sse_dictionaries
    report_data['sse_versions'] = sse_versions

    return report_data


def output_vnv_report(report_info, filename, level):
    """
    This function builds and outputs a V&V report via XLSX file.

    Parameters
    ----------
    report_info: dict
        Dictionary of report information
    filename: str
        Output file name to use to create the report
    level: str
        Level of testing performed (for DVI output)

    Returns
    -------
        None
    """

    # Create a workbook and add a worksheet.
    try:
        workbook = Workbook()
    except:
        msg = f"Failure to open XLSX workbook ({filename}). Can not proceed. Exiting."
        logger.error(msg, exc_info=True)
        raise common.IngeniumLibError(msg)

    worksheet = workbook.active
    worksheet.title = 'Verification Report'

    # Add formats
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))
    title_font = Font(color='FF000000', bold=True)
    title_fill = PatternFill(fill_type='solid', start_color='FFBFBFBF', end_color='FFBFBFBF')
    title_alignment = Alignment(wrap_text=True)
    info_font = Font(color='FF000000', bold=False)
    info_fill = PatternFill(fill_type='solid', start_color='FFFFFFFF', end_color='FFFFFFFF')
    info_alignment = Alignment(wrap_text=True, horizontal='center')

    def write_cell(sheet, row, col, value, font, fill, alignment):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.font = font
        cell.fill = fill
        cell.alignment = alignment
        cell.border = thin_border

    def write_title(sheet, row, col, value):
        write_cell(sheet, row, col, value, title_font, title_fill, title_alignment)

    def write_info(sheet, row, col, value):
        write_cell(sheet, row, col, value, info_font, info_fill, info_alignment)

    # Set Column Widths
    column_widths = [25, 40, 60, 30, 20, 8]
    for idx, width in enumerate(column_widths, 1):
        worksheet.column_dimensions[get_column_letter(idx)].width = width

    row = 1

    # Write the Execution Information
    write_title(worksheet, row, 1, 'Ingenium Execution')
    write_info(worksheet, row, 2, report_info['execution_id'])
    write_title(worksheet, row + 1, 1, 'Time Started')
    write_info(worksheet, row + 1, 2, report_info['time_started'])
    write_title(worksheet, row + 2, 1, 'Time Completed')
    write_info(worksheet, row + 2, 2, report_info['time_completed'])
    write_title(worksheet, row + 3, 1, 'Venue')
    write_info(worksheet, row + 3, 2, report_info['venue_name'])
    write_title(worksheet, row + 4, 1, 'Description')
    write_info(worksheet, row + 4, 2, report_info['description'])
    write_title(worksheet, row + 5, 1, 'Test Conductor(s)')
    write_info(worksheet, row + 5, 2, ','.join(report_info['test_conductors']))

    cur_row = 7
    # Now write the procedure sections
    write_title(worksheet, cur_row, 1, 'Procedures')
    cur_row += 1
    write_title(worksheet, cur_row, 1, 'Procedure ID')
    write_title(worksheet, cur_row, 2, 'Version')
    write_title(worksheet, cur_row, 3, 'Run For Score?')
    write_title(worksheet, cur_row, 4, 'P#')
    write_title(worksheet, cur_row, 5, 'Revision')

    # If there are procedures
    for procedure in report_info.get('used_procedures', []):
            cur_row += 1
            write_info(worksheet, cur_row, 1, procedure['procedure_id'])
            write_info(worksheet, cur_row, 2, procedure['version'])
            write_info(worksheet, cur_row, 3, str(procedure.get('run_for_score', False)))
            write_info(worksheet, cur_row, 4, str(procedure['institutional_id']))
            write_info(worksheet, cur_row, 5, str(procedure['institutional_release_id']))

    cur_row += 2

    # Now write the VI results section
    write_title(worksheet, cur_row, 1, 'VI Type')
    write_title(worksheet, cur_row, 2, 'VI ID')
    write_title(worksheet, cur_row, 3, 'VI Name')
    write_title(worksheet, cur_row, 4, 'Procedure')
    write_title(worksheet, cur_row, 5, 'Step')
    write_title(worksheet, cur_row, 6, 'Status')

    for vi in report_info.get('vis', []):
        cur_row += 1
        worksheet.row_dimensions[cur_row].height = 40
        write_info(worksheet, cur_row, 1, vi['vi_type'])
        write_info(worksheet, cur_row, 2, vi['vi_id'])
        write_info(worksheet, cur_row, 3, vi['vi_name'])
        write_info(worksheet, cur_row, 4, vi['procedure'])
        write_info(worksheet, cur_row, 5, vi['step_number'])
        write_info(worksheet, cur_row, 6, vi['status'])

    # Now build the output in Dictionary "VITS" format
    dvi_sheet = workbook.create_sheet('VITS')

    # Set Column Widths
    dvi_column_widths = [25, 40, 60, 15, 30, 25, 15, 15, 15, 10, 15, 30, 15, 30,
                         25, 25, 25, 25, 25, 25, 25, 25, 25, 25]
    for idx, width in enumerate(dvi_column_widths, 1):
        dvi_sheet.column_dimensions[get_column_letter(idx)].width = width

    # Build the the column headers
    current_row = 1
    dvi_headers = ['Element Type', 'Element ID', 'Element Name', 'Test Level', 'FSW Version',
                   'Dictionary Version', 'SSE Version', 'GDS Version', 'Test Date', 'Status',
                   'Activity Report', 'Test Venue Address', 'MPCS Session ID', 'Test Procedure',
                   'FP Injection Method', 'Notes', 'Anomaly Description', 'PFR ID', 'PFR Link',
                   'PFR Description', 'JIRA Issue', 'JIRA Link', 'JIRA Description', 'Other Issue']
    for col_idx, header in enumerate(dvi_headers, 1):
        write_title(dvi_sheet, current_row, col_idx, header)

    try:
        test_date = report_info['time_completed'].split('T')[0]
    except:
        logger.error(f"Unable to parse time_completed: {report_info['time_completed']}")
        test_date = 'TBD'
    fsw_version = '\n'.join(report_info.get('fsw_versions', []))
    sse_version = '\n'.join(report_info.get('sse_versions', []))
    gds_version = '\n'.join(report_info.get('gds_versions', []))
    dictionary_version = '\n'.join(report_info.get('fsw_dictionaries', []))
    test_venue_address = '\n'.join(report_info.get('test_venue_addresses', []))
    ampcs_sessions = ','.join(report_info.get('gds_sessions', []))
    activity_reports = ','.join(report_info.get('activity_reports', []))

    test_level = TEST_LEVELS.get(level)
    if test_level is None:
        logger.error('Unable to determine test level. Defaulting to systems.')
        test_level = 'System I&T'

    for vi in report_info.get('vis', []):
        current_row = current_row + 1
        write_info(dvi_sheet, current_row, 1, vi['vi_type'])
        write_info(dvi_sheet, current_row, 2, vi['vi_id'])
        write_info(dvi_sheet, current_row, 3, vi['vi_name'])
        write_info(dvi_sheet, current_row, 4, test_level)
        write_info(dvi_sheet, current_row, 5, fsw_version)
        write_info(dvi_sheet, current_row, 6, dictionary_version)
        write_info(dvi_sheet, current_row, 7, sse_version)
        write_info(dvi_sheet, current_row, 8, gds_version)
        write_info(dvi_sheet, current_row, 9, test_date)
        write_info(dvi_sheet, current_row, 10, vi['status'])
        write_info(dvi_sheet, current_row, 11, activity_reports)
        write_info(dvi_sheet, current_row, 12, test_venue_address)
        write_info(dvi_sheet, current_row, 13, ampcs_sessions)
        write_info(dvi_sheet, current_row, 14, vi['inst_procedure'])

    try:
        workbook.save(filename)
        workbook.close()
    except:
        msg = 'Unable to save V&V Report.'
        logger.error(msg, exc_info=True)
        raise common.IngeniumLibError(msg)

    return

##################################################### Main ###########################################################


def main(args=[]):
    """
    This is the main function of the Ingenium VNV Report Generator

    Parameters
    ----------
    args
        Array of input arguments. Used when this function is called from another Python module.

    Returns
    -------
        None
    """

    # Gets initial input
    inputs = get_input(args)

    # if ssl_ca_bundle is specified, it will take precedence over ignore_ssl_error
    if inputs.ssl_ca_bundle is not None:
        common.set_ssl_verify(inputs.ssl_ca_bundle)
    else:
        common.set_ssl_verify(not inputs.ignore_ssl_error)
        if not common.get_ssl_verify():
            # To suppress SSL warnings
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    logger.debug(f"Ingenium.ssl_verify:{common.get_ssl_verify()}")

    if inputs.username:
        username = inputs.username
    else:
        username = getpass.getuser()

    if inputs.rsa:
        pw_prompt = f"Enter RSA Passcode for {username}:"
    else:
        pw_prompt = f"Enter LDAP Password for {username}:"

    # Login to Ingenium
    login = common.authenticate(inputs.ingenium_server, username=username,
                                password=getpass.getpass(pw_prompt), force=True,
                                rsa=inputs.rsa)
    if not login:
        msg = 'Failure to Login - Can not proceed. Exiting.'
        logger.error(msg)
        raise common.IngeniumLibError(msg)

    # Query the As-Run
    execution_info, execution_data = get_ingenium_execution_elements(inputs.ingenium_server,
                                                                       inputs.execution_id)

    # Build Information to create the report
    report_info = build_report_data(execution_data, execution_info)

    # Write the Report to a File
    output_vnv_report(report_info, inputs.output_file, inputs.level)


if __name__ == '__main__':
    try:
        main()
    except common.IngeniumLibError:
        logger.error("Ingenium V&V Report encountered an error and needs to exit. Please check the log messages for the source of the error.")
